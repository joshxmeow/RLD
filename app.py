import json
import hashlib
import math
import os
import re
import secrets
import socket
import sqlite3
import tempfile
from datetime import datetime, timedelta, timezone
from contextlib import contextmanager
from functools import wraps
from pathlib import Path

from flask import Flask, abort, jsonify, request, send_from_directory, session
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.worksheet.page import PageMargins
from werkzeug.security import check_password_hash, generate_password_hash

ROOT = Path(__file__).resolve().parent
STORAGE_CONFIG_PATH = ROOT / "storage-config.json"
ADMIN_USERNAME = "rld"
ADMIN_PASSWORD = "rld2026"
SERVER_PORT = 8443
USERNAME_PATTERN = re.compile(r"^[a-z0-9._-]{3,30}$")
PUBLIC_FILES = {
    "index.html", "style.css", "script.js", "manifest.webmanifest",
    "service-worker.js", "rld_logo.png", "header.png", "signature1.png", "signature2.png",
}


def resolve_data_dir(config_path=STORAGE_CONFIG_PATH):
    configured_path = os.environ.get("RLD_DATA_DIR")
    if not configured_path:
        config_path = Path(config_path)
        if config_path.exists():
            try:
                configuration = json.loads(config_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as error:
                raise RuntimeError(f"Could not read storage configuration at {config_path}: {error}") from error
            configured_path = configuration.get("data_dir") if isinstance(configuration, dict) else None
            if not isinstance(configured_path, str) or not configured_path.strip():
                raise RuntimeError(f"Storage configuration at {config_path} has no valid data_dir.")
    data_dir = Path(configured_path).expanduser() if configured_path else ROOT / "instance"
    return data_dir.resolve()


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")

def format_estimate_date(value):
    try:
        parsed_date = datetime.strptime(value, "%Y-%m-%d")
    except (TypeError, ValueError):
        return str(value or "")
    month_names = (
        "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
        "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
    )
    return f"{month_names[parsed_date.month - 1]} {parsed_date.day}, {parsed_date.year}"


def connect_db(database_path):
    connection = sqlite3.connect(database_path, timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


@contextmanager
def database_connection(database_path):
    connection = connect_db(database_path)
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def init_db(database_path):
    Path(database_path).parent.mkdir(parents=True, exist_ok=True)
    with database_connection(database_path) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                username TEXT PRIMARY KEY COLLATE NOCASE,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK (role IN ('admin', 'user')),
                created_at TEXT NOT NULL,
                password_version INTEGER NOT NULL DEFAULT 1
            )
            """
        )
        user_columns = {
            row["name"] for row in connection.execute("PRAGMA table_info(users)").fetchall()
        }
        if "password_version" not in user_columns:
            connection.execute(
                "ALTER TABLE users ADD COLUMN password_version INTEGER NOT NULL DEFAULT 1"
            )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS estimates (
                id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                server_updated_at TEXT NOT NULL,
                client_updated_at TEXT NOT NULL,
                username TEXT NOT NULL,
                data_json TEXT NOT NULL,
                deleted INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS units (
                id TEXT PRIMARY KEY,
                maker TEXT NOT NULL,
                unit TEXT NOT NULL,
                transmission TEXT NOT NULL,
                plate_no TEXT NOT NULL,
                year_model INTEGER NOT NULL,
                fuel_type TEXT NOT NULL,
                selling_price REAL NOT NULL CHECK (selling_price >= 0),
                available INTEGER NOT NULL DEFAULT 1 CHECK (available IN (0, 1))
            )
            """
        )
        admin = connection.execute(
            "SELECT username FROM users WHERE username = ?", (ADMIN_USERNAME,)
        ).fetchone()
        if admin is None:
            connection.execute(
                "INSERT INTO users (username, password_hash, role, created_at) VALUES (?, ?, ?, ?)",
                (ADMIN_USERNAME, generate_password_hash(ADMIN_PASSWORD), "admin", utc_now()),
            )


def create_app(database_path=None, secret_key=None):
    data_dir = Path(database_path).parent if database_path is not None else resolve_data_dir()
    database_path = Path(database_path) if database_path is not None else data_dir / "rld.sqlite3"
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=secret_key or load_secret_key(data_dir),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SECURE=True,
        SESSION_COOKIE_SAMESITE="Lax",
        PERMANENT_SESSION_LIFETIME=timedelta(days=30),
        MAX_CONTENT_LENGTH=2 * 1024 * 1024,
        DATABASE=str(database_path),
    )
    init_db(app.config["DATABASE"])

    def current_user():
        username = session.get("username")
        if not username:
            return None
        with database_connection(app.config["DATABASE"]) as connection:
            user = connection.execute(
                "SELECT username, role, password_version FROM users WHERE username = ?", (username,)
            ).fetchone()
        if user is None or session.get("auth_version") != user["password_version"]:
            return None
        return user

    def login_required(handler):
        @wraps(handler)
        def wrapped(*args, **kwargs):
            if current_user() is None:
                return jsonify(error="Please sign in to the local server again.", code="session_expired"), 401
            return handler(*args, **kwargs)
        return wrapped

    def admin_required(handler):
        @wraps(handler)
        @login_required
        def wrapped(*args, **kwargs):
            user = current_user()
            if user["role"] != "admin":
                return jsonify(error="Administrator access is required."), 403
            return handler(*args, **kwargs)
        return wrapped

    def require_same_origin():
        origin = request.headers.get("Origin")
        if not origin:
            return
        expected = request.host_url.rstrip("/")
        if origin.rstrip("/") != expected:
            abort(403, description="Requests from another website are not allowed.")

    def estimate_record(row):
        return {
            "id": row["id"],
            "savedAt": row["created_at"],
            "updatedAt": row["client_updated_at"],
            "serverUpdatedAt": row["server_updated_at"],
            "username": row["username"],
            "data": json.loads(row["data_json"]),
            "deleted": bool(row["deleted"]),
        }

    def unit_record(row):
        return {
            "id": row["id"],
            "maker": row["maker"],
            "unit": row["unit"],
            "transmission": row["transmission"],
            "plateNo": row["plate_no"],
            "yearModel": row["year_model"],
            "fuelType": row["fuel_type"],
            "sellingPrice": row["selling_price"],
            "available": bool(row["available"]),
        }

    def validated_unit(payload):
        if not isinstance(payload, dict):
            return None, "Enter valid unit details."
        makers = ("Toyota", "Honda", "Isuzu", "Mitsubishi", "Hyundai", "Ford", "Volkswagen", "Nissan")
        maker = payload.get("maker")
        text_fields = {
            "unit": ("unit", 100),
            "transmission": ("transmission", 40),
            "plateNo": ("plateNo", 20),
            "fuelType": ("fuelType", 40),
        }
        optional_fields = {"plateNo"}
        if maker not in makers:
            return None, "Choose a maker from the list."
        values = {}
        for field, (key, limit) in text_fields.items():
            value = payload.get(key, "" if field in optional_fields else None)
            if (
                not isinstance(value, str)
                or (field not in optional_fields and not value.strip())
                or len(value.strip()) > limit
            ):
                return None, f"Enter a valid {field.replace('No', ' number').replace('Type', ' type')}."
            values[field] = value.strip()
        year_model = payload.get("yearModel")
        if isinstance(year_model, bool):
            return None, "Enter a valid year model."
        if isinstance(year_model, float) and not year_model.is_integer():
            return None, "Enter a valid year model."
        try:
            year_model = int(year_model)
        except (TypeError, ValueError, OverflowError):
            return None, "Enter a valid year model."
        if not 1900 <= year_model <= 2100:
            return None, "Enter a valid year model."
        selling_price = payload.get("sellingPrice")
        if isinstance(selling_price, bool):
            return None, "Enter a valid selling price."
        try:
            selling_price = float(selling_price)
        except (TypeError, ValueError, OverflowError):
            return None, "Enter a valid selling price."
        if not 0 <= selling_price <= 9999999999.99 or not math.isfinite(selling_price):
            return None, "Enter a valid selling price."
        available = payload.get("available", True)
        if not isinstance(available, bool):
            return None, "Choose a valid availability status."
        return {
            "maker": maker,
            "unit": values["unit"],
            "transmission": values["transmission"],
            "plateNo": values["plateNo"],
            "yearModel": year_model,
            "fuelType": values["fuelType"],
            "sellingPrice": round(selling_price, 2),
            "available": available,
        }, None

    @app.get("/")
    def home():
        return send_from_directory(ROOT, "index.html")

    @app.get("/<path:filename>")
    def app_files(filename):
        if filename not in PUBLIC_FILES:
            abort(404)
        return send_from_directory(ROOT, filename)

    @app.get("/api/health")
    def health():
        return jsonify(status="ok", serverTime=utc_now())

    @app.get("/api/units")
    @login_required
    def list_units():
        with database_connection(app.config["DATABASE"]) as connection:
            rows = connection.execute(
                "SELECT * FROM units ORDER BY maker COLLATE NOCASE, unit COLLATE NOCASE, year_model DESC"
            ).fetchall()
        return jsonify(units=[unit_record(row) for row in rows])

    @app.post("/api/units")
    @login_required
    def create_unit():
        require_same_origin()
        unit, error = validated_unit(request.get_json(silent=True))
        if error:
            return jsonify(error=error), 400
        unit_id = secrets.token_urlsafe(12)
        with database_connection(app.config["DATABASE"]) as connection:
            connection.execute(
                """
                INSERT INTO units
                    (id, maker, unit, transmission, plate_no, year_model, fuel_type, selling_price, available)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    unit_id, unit["maker"], unit["unit"], unit["transmission"], unit["plateNo"],
                    unit["yearModel"], unit["fuelType"], unit["sellingPrice"], int(unit["available"]),
                ),
            )
            saved = connection.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
        return jsonify(unit=unit_record(saved)), 201

    @app.put("/api/units/<unit_id>")
    @login_required
    def update_unit(unit_id):
        require_same_origin()
        unit, error = validated_unit(request.get_json(silent=True))
        if error:
            return jsonify(error=error), 400
        available = unit["available"]
        with database_connection(app.config["DATABASE"]) as connection:
            result = connection.execute(
                """
                UPDATE units SET maker = ?, unit = ?, transmission = ?, plate_no = ?,
                    year_model = ?, fuel_type = ?, selling_price = ?, available = ?
                WHERE id = ?
                """,
                (
                    unit["maker"], unit["unit"], unit["transmission"], unit["plateNo"],
                    unit["yearModel"], unit["fuelType"], unit["sellingPrice"], int(available), unit_id,
                ),
            )
            if result.rowcount == 0:
                return jsonify(error="That unit no longer exists."), 404
            saved = connection.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
        return jsonify(unit=unit_record(saved))

    @app.delete("/api/units/<unit_id>")
    @login_required
    def delete_unit(unit_id):
        require_same_origin()
        with database_connection(app.config["DATABASE"]) as connection:
            result = connection.execute("DELETE FROM units WHERE id = ?", (unit_id,))
        if result.rowcount == 0:
            return jsonify(error="That unit no longer exists."), 404
        return jsonify(id=unit_id, status="deleted")

    @app.post("/api/login")
    def login():
        require_same_origin()
        payload = request.get_json(silent=True) or {}
        username = str(payload.get("username", "")).strip().lower()
        password = payload.get("password", "")
        if not isinstance(password, str) or not username:
            return jsonify(error="Enter a username and password.", code="missing_credentials"), 400
        with database_connection(app.config["DATABASE"]) as connection:
            user = connection.execute(
                "SELECT username, password_hash, role, password_version FROM users WHERE username = ?",
                (username,),
            ).fetchone()
        if user is None:
            return jsonify(error="Username not found on this local server.", code="username_not_found"), 404
        if not check_password_hash(user["password_hash"], password):
            return jsonify(error="Incorrect password for this username.", code="wrong_password"), 401
        session.clear()
        session.permanent = bool(payload.get("remember"))
        session["username"] = user["username"]
        session["auth_version"] = user["password_version"]
        return jsonify(username=user["username"], role=user["role"])

    @app.post("/api/logout")
    def logout():
        require_same_origin()
        session.clear()
        return jsonify(status="ok")

    @app.get("/api/session")
    @login_required
    def session_info():
        user = current_user()
        return jsonify(username=user["username"], role=user["role"])

    @app.put("/api/admin/password")
    @admin_required
    def change_admin_password():
        require_same_origin()
        payload = request.get_json(silent=True) or {}
        current_password = payload.get("currentPassword", "")
        new_password = payload.get("newPassword", "")
        if not isinstance(current_password, str) or not isinstance(new_password, str):
            return jsonify(error="Enter the current and new passwords."), 400
        if len(new_password) < 8:
            return jsonify(error="Passwords must be at least 8 characters."), 400
        with database_connection(app.config["DATABASE"]) as connection:
            user = connection.execute(
                "SELECT password_hash FROM users WHERE username = ?", (ADMIN_USERNAME,)
            ).fetchone()
            if (
                user is None
                or not check_password_hash(user["password_hash"], current_password)
            ):
                return jsonify(error="The current administrator password is incorrect."), 401
            connection.execute(
                "UPDATE users SET password_hash = ?, password_version = password_version + 1 WHERE username = ?",
                (generate_password_hash(new_password), ADMIN_USERNAME),
            )
        session.clear()
        return jsonify(status="updated")

    @app.get("/api/users")
    @admin_required
    def list_users():
        with database_connection(app.config["DATABASE"]) as connection:
            users = connection.execute(
                "SELECT username, role, created_at FROM users ORDER BY username COLLATE NOCASE"
            ).fetchall()
        return jsonify(users=[dict(user) for user in users])

    @app.post("/api/users")
    @admin_required
    def create_user():
        require_same_origin()
        payload = request.get_json(silent=True) or {}
        username = str(payload.get("username", "")).strip().lower()
        password = payload.get("password", "")
        if not USERNAME_PATTERN.fullmatch(username) or username == ADMIN_USERNAME:
            return jsonify(error="Use a unique 3-30 character username (letters, numbers, dot, underscore, hyphen)."), 400
        if not isinstance(password, str) or len(password) < 8:
            return jsonify(error="Passwords must be at least 8 characters."), 400
        try:
            with database_connection(app.config["DATABASE"]) as connection:
                connection.execute(
                    "INSERT INTO users (username, password_hash, role, created_at) VALUES (?, ?, 'user', ?)",
                    (username, generate_password_hash(password), utc_now()),
                )
        except sqlite3.IntegrityError:
            return jsonify(error="That username already exists."), 409
        return jsonify(username=username, role="user"), 201

    @app.put("/api/users/<username>")
    @admin_required
    def change_user_password(username):
        require_same_origin()
        if username.lower() == ADMIN_USERNAME:
            return jsonify(error="Change the administrator password directly on the Windows server."), 400
        payload = request.get_json(silent=True) or {}
        password = payload.get("password", "")
        if not isinstance(password, str) or len(password) < 8:
            return jsonify(error="Passwords must be at least 8 characters."), 400
        with database_connection(app.config["DATABASE"]) as connection:
            result = connection.execute(
                "UPDATE users SET password_hash = ?, password_version = password_version + 1 WHERE username = ?",
                (generate_password_hash(password), username),
            )
        if result.rowcount == 0:
            return jsonify(error="That account no longer exists."), 404
        return jsonify(username=username, status="updated")

    @app.delete("/api/users/<username>")
    @admin_required
    def delete_user(username):
        require_same_origin()
        if username.lower() == ADMIN_USERNAME:
            return jsonify(error="The administrator account cannot be deleted."), 400
        with database_connection(app.config["DATABASE"]) as connection:
            result = connection.execute("DELETE FROM users WHERE username = ?", (username,))
        if result.rowcount == 0:
            return jsonify(error="That account no longer exists."), 404
        return jsonify(username=username, status="deleted")

    @app.get("/api/estimates")
    @login_required
    def list_estimates():
        with database_connection(app.config["DATABASE"]) as connection:
            rows = connection.execute(
                "SELECT * FROM estimates ORDER BY server_updated_at DESC"
            ).fetchall()
        return jsonify(estimates=[estimate_record(row) for row in rows])

    @app.put("/api/estimates/<estimate_id>")
    @login_required
    def save_estimate(estimate_id):
        require_same_origin()
        payload = request.get_json(silent=True) or {}
        data = payload.get("data")
        if not isinstance(data, dict) or payload.get("id") != estimate_id:
            return jsonify(error="Invalid estimate data."), 400
        client_updated_at = payload.get("updatedAt")
        if not isinstance(client_updated_at, str) or len(client_updated_at) > 40:
            return jsonify(error="The estimate timestamp is invalid."), 400
        try:
            parsed_updated_at = datetime.fromisoformat(client_updated_at.replace("Z", "+00:00"))
            if parsed_updated_at.tzinfo is None:
                raise ValueError("Timezone is required.")
            client_updated_at = parsed_updated_at.astimezone(timezone.utc).isoformat(
                timespec="milliseconds"
            ).replace("+00:00", "Z")
        except ValueError:
            return jsonify(error="The estimate timestamp is invalid."), 400
        deleted = payload.get("deleted", False)
        if not isinstance(deleted, bool):
            return jsonify(error="The estimate deletion flag is invalid."), 400
        now = utc_now()
        user = current_user()
        with database_connection(app.config["DATABASE"]) as connection:
            existing = connection.execute(
                "SELECT * FROM estimates WHERE id = ?", (estimate_id,)
            ).fetchone()
            if existing and existing["client_updated_at"] > client_updated_at:
                return jsonify(accepted=False, estimate=estimate_record(existing))
            created_at = existing["created_at"] if existing else now
            connection.execute(
                """
                INSERT INTO estimates
                    (id, created_at, server_updated_at, client_updated_at, username, data_json, deleted)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(id) DO UPDATE SET
                    server_updated_at = excluded.server_updated_at,
                    client_updated_at = excluded.client_updated_at,
                    username = excluded.username,
                    data_json = excluded.data_json,
                    deleted = excluded.deleted
                """,
                (
                    estimate_id, created_at, now, client_updated_at, user["username"],
                    json.dumps(data, ensure_ascii=False), int(deleted),
                ),
            )
            saved = connection.execute(
                "SELECT * FROM estimates WHERE id = ?", (estimate_id,)
            ).fetchone()
        if not deleted:
            try:
                save_estimate_workbook(estimate_record(saved), data_dir)
            except (OSError, ValueError) as error:
                app.logger.exception("Could not write the estimate workbook for %s", estimate_id)
                return jsonify(
                    error=f"Estimate saved on the server, but its Excel file could not be written: {error}"
                ), 500
        return jsonify(accepted=True, estimate=estimate_record(saved))

    return app


def estimate_filename_component(value, fallback):
    component = str(value or "").strip().upper()
    component = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", component).rstrip(" .")
    component = component[:80].rstrip(" .")
    if not component:
        component = fallback
    if component.split(".", 1)[0] in {
        "CON", "PRN", "AUX", "NUL",
        *(f"COM{number}" for number in range(1, 10)),
        *(f"LPT{number}" for number in range(1, 10)),
    }:
        component = f"_{component}"
    return component


def estimate_workbook_filename(data):
    organization = estimate_filename_component(data.get("orgPrefix"), "PRIVATE")
    make_model = estimate_filename_component(data.get("makeModel"), "UNKNOWN VEHICLE")
    plate_number = estimate_filename_component(data.get("plateNo"), "NO PLATE")
    return f"{organization}_{make_model}_{plate_number}.xlsx"


def estimate_work_rows(data):
    rows = []
    number = 1
    works = data.get("works", [])
    if not isinstance(works, list):
        return rows
    for item in works:
        if not isinstance(item, dict):
            continue
        description = str(item.get("desc", "")).strip()
        if not description:
            continue
        if description.startswith("*"):
            rows.append((description.upper(), item.get("labor", ""), item.get("parts", "")))
            number = 1
            continue
        description = re.sub(r"^\d+\.\s*", "", description)
        rows.append((f"{number}. {description.upper()}", item.get("labor", ""), item.get("parts", "")))
        number += 1
    return rows


def estimate_work_rows_with_sections(data):
    rows = []
    sections = []
    number = 1
    section_start = 0
    is_parts_section = False
    works = data.get("works", [])
    if not isinstance(works, list):
        return rows, sections

    def finish_section(end):
        nonlocal section_start
        if end > section_start:
            sections.append((section_start, end, is_parts_section))
        section_start = end

    for item in works:
        if not isinstance(item, dict):
            continue
        description = str(item.get("desc", "")).strip()
        if not description:
            continue
        if description.startswith("*"):
            finish_section(len(rows))
            rows.append((description.upper(), item.get("labor", ""), item.get("parts", "")))
            section_start = len(rows)
            is_parts_section = description.upper().replace(" ", "") == "*PARTS"
            number = 1
            continue
        description = re.sub(r"^\d+\.\s*", "", description)
        rows.append((f"{number}. {description.upper()}", item.get("labor", ""), item.get("parts", "")))
        number += 1
    finish_section(len(rows))
    return rows, sections


def clean_excel_amount(value):
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value or ""))
    try:
        return float(cleaned) if cleaned else 0
    except ValueError:
        return 0


def excel_has_value(value):
    return value is not None and str(value).strip() != ""


def apply_merged_border(sheet, first_row, last_row, column, side):
    for row in range(first_row, last_row + 1):
        sheet.cell(row, column).border = Border(
            left=side,
            right=side,
            top=side if row == first_row else Side(style=None),
            bottom=side if row == last_row else Side(style=None),
        )


def save_estimate_workbook(record, data_dir):
    data = record["data"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Job Estimate"
    sheet.sheet_view.showGridLines = False
    for column in ("A", "B", "C", "D"):
        sheet.column_dimensions[column].width = 25
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    # Excel paper code 14 is Folio, matching the browser's 8.5 x 13-inch print page.
    sheet.page_setup.paperSize = "14"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.3, bottom=0.3, header=0, footer=0)
    sheet.print_options.horizontalCentered = True

    thin_black = Side(style="thin", color="000000")
    grid = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    bold = Font(name="Arial", size=9, bold=True)
    normal = Font(name="Arial", size=8, bold=True)

    banner_path = ROOT / "header.png"
    if banner_path.exists():
        banner = ExcelImage(str(banner_path))
        original_width, original_height = banner.width, banner.height
        scale = min(700 / original_width, 80 / original_height)
        banner.width = round(original_width * scale)
        banner.height = round(original_height * scale)
        banner.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=0,
                row=0,
                colOff=pixels_to_EMU(10),
                rowOff=pixels_to_EMU(9),
            ),
            ext=XDRPositiveSize2D(
                cx=pixels_to_EMU(banner.width),
                cy=pixels_to_EMU(banner.height),
            ),
        )
        sheet.add_image(banner)
    sheet.row_dimensions[1].height = 60
    for row, value, size in (
        (2, "RLD AUTOWORX AND TOWING SERVICES", 12),
        (3, "Brgy. Buli, Taal, Batangas", 9),
        (4, "Contact No: 09296462336/09267433265/(043)321-8477", 9),
        (5, "Email: rldcarpaintingservices@yahoo.com", 9),
        (6, "TIN NO: 287-920-998-000", 9),
    ):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = sheet.cell(row, 1, value)
        cell.font = Font(name="Arial", size=size, bold=True)
        cell.alignment = centered

    sheet.merge_cells("A8:D8")
    sheet["A8"] = "JOB ESTIMATE"
    sheet["A8"].font = Font(name="Arial", size=12, bold=True)
    sheet["A8"].alignment = centered
    sheet["A9"] = f"DATE: {format_estimate_date(data.get('date'))}"
    sheet.merge_cells("A9:D9")
    sheet["A9"].font = bold
    sheet["A9"].alignment = right

    sheet.merge_cells("A11:B11")
    sheet.merge_cells("C11:D11")
    sheet["A11"] = "CLIENT"
    sheet["C11"] = "VEHICLE"
    for column in range(1, 5):
        sheet.cell(11, column).font = bold
        sheet.cell(11, column).alignment = centered
        sheet.cell(11, column).border = grid

    details = [
        (f"NAME: {data.get('clientName', '')}", f"MAKE & MODEL: {data.get('makeModel', '')}", f"PLATE NO.: {data.get('plateNo', '')}"),
        (f"ADDRESS: {data.get('address', '')}", f"CHASSIS NO: {data.get('chassisNo', '')}", f"ENGINE NO.: {data.get('engineNo', '')}"),
        (f"CONTACT NUMBER: {data.get('contact', '')}", f"COLOR: {data.get('color', '')}", f"INSURANCE.: {data.get('insurance', '')}"),
    ]
    detail_column_widths = (50, 25, 25)
    for row_index, values in enumerate(details, start=12):
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
        for column, value in ((1, values[0]), (3, values[1]), (4, values[2])):
            cell = sheet.cell(row_index, column, value)
            cell.font = normal
            cell.alignment = left
        for column in range(1, 5):
            sheet.cell(row_index, column).border = grid
        lines = max(
            max(1, (len(str(value)) + width - 1) // width)
            for value, width in zip(values, detail_column_widths)
        )
        sheet.row_dimensions[row_index].height = max(16, lines * 12)

    work_header_row = 16
    sheet.merge_cells(start_row=work_header_row, start_column=1, end_row=work_header_row, end_column=2)
    for column, value in ((1, "SCOPE OF WORKS"), (3, "COST OF LABOR"), (4, "COST OF PARTS")):
        cell = sheet.cell(work_header_row, column, value)
        cell.font = bold
        cell.alignment = centered
    for column in range(1, 5):
        sheet.cell(work_header_row, column).border = grid
    row_index = work_header_row + 1
    work_rows, work_sections = estimate_work_rows_with_sections(data)
    work_row_start = row_index
    for description, labor, parts in work_rows:
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
        for column, value in ((1, description), (3, labor), (4, parts)):
            cell = sheet.cell(row_index, column, value)
            cell.font = bold if description.startswith("*") else normal
            cell.alignment = left if column == 1 else right
            if column in (3, 4) and isinstance(value, (int, float)):
                cell.number_format = '"PHP " #,##0.00'
        for column in range(1, 5):
            sheet.cell(row_index, column).border = grid
        sheet.row_dimensions[row_index].height = 16
        row_index += 1

    if data.get("noteHiddenDamage") is True:
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
        note = sheet.cell(row_index, 1, "NOTE: HIDDEN DAMAGE ARE NOT INCLUDED")
        note.font = Font(name="Arial", size=9, bold=True, color="C00000")
        note.alignment = left
        for column in range(1, 5):
            sheet.cell(row_index, column).border = grid
        sheet.row_dimensions[row_index].height = 16
        row_index += 1

    bundle_all_prices = data.get("bundleAllPrices", False) is True
    bundle_labor = data.get("bundleLabor", True) is True
    for section_start, section_end, is_parts_section in work_sections:
        if section_end - section_start < 2:
            continue
        first_sheet_row = work_row_start + section_start
        last_sheet_row = work_row_start + section_end - 1
        if bundle_all_prices:
            for column, cost_index in ((3, 1), (4, 2)):
                section_values = [work_rows[index][cost_index] for index in range(section_start, section_end)]
                total = sum(clean_excel_amount(value) for value in section_values)
                if total == 0:
                    continue
                sheet.cell(first_sheet_row, column, total)
                sheet.cell(first_sheet_row, column).number_format = '"PHP " #,##0.00'
                sheet.merge_cells(
                    start_row=first_sheet_row, start_column=column,
                    end_row=last_sheet_row, end_column=column
                )
                apply_merged_border(sheet, first_sheet_row, last_sheet_row, column, thin_black)
        elif bundle_labor and not is_parts_section:
            labor_values = [
                work_rows[index][1] for index in range(section_start, section_end)
                if excel_has_value(work_rows[index][1])
            ]
            if len(labor_values) == 1:
                sheet.cell(first_sheet_row, 3, clean_excel_amount(labor_values[0]))
                sheet.cell(first_sheet_row, 3).number_format = '"PHP " #,##0.00'
                sheet.merge_cells(
                    start_row=first_sheet_row, start_column=3,
                    end_row=last_sheet_row, end_column=3
                )
                apply_merged_border(sheet, first_sheet_row, last_sheet_row, 3, thin_black)

    totals = [
        ("TOTAL LABOR/MATERIALS/PARTS:", data.get("totLabor", 0), data.get("totParts", 0)),
        ("DOWNPAYMENT:", data.get("downpayment", 0), None),
        ("TOTAL AMOUNT DUE:", data.get("totDue", 0), None),
    ]
    for label, first_value, second_value in totals:
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
        sheet.cell(row_index, 1, label)
        for column, value in ((3, first_value), (4, second_value)):
            if value is not None:
                cell = sheet.cell(row_index, column, value)
                if isinstance(value, (int, float)):
                    cell.number_format = '"PHP " #,##0.00'
        if second_value is None:
            sheet.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=4)
        for column in range(1, 5):
            sheet.cell(row_index, column).font = bold
            sheet.cell(row_index, column).alignment = right
            sheet.cell(row_index, column).border = grid
        row_index += 1

    row_index += 1
    signatures = []
    for image_name, column, label in (
        ("signature1.png", 1, "PREPARED BY: BEVERLY BALATBAT"),
        ("signature2.png", 3, "APPROVED BY: ROLDAN DUEÑAS"),
    ):
        image_path = ROOT / image_name
        if image_path.exists():
            signature = ExcelImage(str(image_path))
            original_width, original_height = signature.width, signature.height
            signature.height = 65
            signature.width = round(original_width * signature.height / original_height)
            center_offset = pixels_to_EMU(180) - pixels_to_EMU(signature.width) // 2
            signature.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=column - 1,
                    row=row_index - 1,
                    colOff=center_offset,
                    rowOff=pixels_to_EMU(4),
                ),
                ext=XDRPositiveSize2D(
                    cx=pixels_to_EMU(signature.width),
                    cy=pixels_to_EMU(signature.height),
                ),
            )
            sheet.add_image(signature)
            signatures.append((column, label))
    sheet.row_dimensions[row_index].height = 56
    for column, label in signatures:
        cell = sheet.cell(row_index + 1, column, label)
        cell.value = label
        cell.font = bold
        cell.alignment = centered
        sheet.merge_cells(
            start_row=row_index + 1, start_column=column,
            end_row=row_index + 1, end_column=column + 1
        )
        for border_column in range(column, column + 2):
            sheet.cell(row_index + 1, border_column).border = Border(bottom=thin_black)
    row_index += 2

    terms = (
        "I HEREBY AUTHORIZE AND AGREE TO PAY FOR THE ABOVE REPAIRS AND WORK TO BE DONE ON MY VEHICLE "
        "INCLUDING ALL PARTS AND MATERIALS NECESSARY TO PERFORM THE SAME. I ALSO AGREE TO THE TERMS AND "
        "CONDITION STATED BELOW:\n"
        "1. RLD AUTOWORX & TOWING SERVICES will not be held responsible for any losses or damages to the vehicle "
        "and its accessories or any articles of value left in the vehicle in case of theft, fire, accident or any "
        "other cause beyond the auto center's control.\n"
        "2. Customer acknowledges an expenses mechanic lien to secure the amount of repairs thereto.\n"
        "3. Company employees may operate vehicle for purposes of testing inspection and /or deliver at customer's risk.\n"
        "4. No claims for damages to vehicle under this repair order will be considered unless received by the company "
        "within three (3) days after the vehicle has been delivered.\n"
        "5. Customers agrees to pay interest at the rate of 3% per month on all accounts not paid when due.\n"
        "6. RLD AUTOWORX & TOWING SERVICES will charge a daily storage fee of P300.00 on the vehicle from the day "
        "enter on the repair shop if not yet started of repair of vehicle.\n"
        "7. RLD AUTOWORX & TOWING SERVICES to charge a daily storage fee of P300.00 on the vehicle not claimed on "
        "the date of upon completion of repair.\n"
        "8. Hidden defects are not included in this evaluation\n"
        "9. Prices are subject to change without prior notice."
    )
    sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=4)
    terms_cell = sheet.cell(row_index, 1, terms)
    terms_cell.font = Font(name="Arial", size=5.5)
    terms_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    sheet.row_dimensions[row_index].height = 175
    sheet.print_area = f"A1:D{row_index}"

    estimate_id = str(record["id"])
    estimate_dir = Path(data_dir) / "estimates" / hashlib.sha256(estimate_id.encode("utf-8")).hexdigest()
    estimate_dir.mkdir(parents=True, exist_ok=True)
    destination = estimate_dir / estimate_workbook_filename(data)
    with tempfile.NamedTemporaryFile(dir=estimate_dir, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)
    for old_file in estimate_dir.glob("*.xlsx"):
        if old_file != destination:
            old_file.unlink()
    return destination


def load_secret_key(data_dir):
    data_dir.mkdir(parents=True, exist_ok=True)
    key_path = data_dir / "session-secret.key"
    if key_path.exists():
        return key_path.read_text(encoding="ascii").strip()
    secret = secrets.token_hex(32)
    key_path.write_text(secret, encoding="ascii")
    return secret


def local_ipv4():
    probe = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        probe.connect(("192.0.2.1", 1))
        return probe.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        probe.close()


if __name__ == "__main__":
    from cert_setup import ensure_certificates

    data_dir = resolve_data_dir()
    lan_ip = local_ipv4()
    cert_path, key_path, ca_path = ensure_certificates(data_dir / "certs", lan_ip)
    app = create_app()
    print(f"RLD local server: https://{lan_ip}:{SERVER_PORT}")
    print(f"Install this CA certificate on phones: {ca_path}")
    print(f"Estimate database and server keys are stored in: {data_dir}")
    print("Keep this data folder private. Press Ctrl+C to stop.")
    app.run(host="0.0.0.0", port=SERVER_PORT, ssl_context=(str(cert_path), str(key_path)), threaded=True)
