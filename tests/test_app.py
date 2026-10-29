from io import BytesIO
import tempfile
import unittest
import sqlite3
import zipfile
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image

from app import ADMIN_PASSWORD, ADMIN_USERNAME, create_app, resolve_data_dir


class LocalServerTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        database_path = Path(self.temporary_directory.name) / "test.sqlite3"
        self.app = create_app(database_path=database_path, secret_key="test-secret")
        self.app.config.update(TESTING=True)
        self.client = self.app.test_client()
        self.base_url = "https://localhost"

    def tearDown(self):
        self.temporary_directory.cleanup()

    def login(self, client=None, username=ADMIN_USERNAME, password=ADMIN_PASSWORD):
        return (client or self.client).post(
            "/api/login",
            base_url=self.base_url,
            json={"username": username, "password": password},
        )

    def test_data_folder_can_be_configured_by_file_or_environment(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_path = root / "storage-config.json"
            configured_dir = root / "chosen-data"
            config_path.write_text(
                '{"data_dir": "' + str(configured_dir).replace("\\", "\\\\") + '"}',
                encoding="utf-8",
            )
            self.assertEqual(resolve_data_dir(config_path), configured_dir.resolve())

            environment_dir = root / "environment-data"
            with patch.dict("os.environ", {"RLD_DATA_DIR": str(environment_dir)}):
                self.assertEqual(resolve_data_dir(config_path), environment_dir.resolve())
                configured_app = create_app(database_path=None, secret_key="test-secret")
                self.assertEqual(
                    configured_app.config["DATABASE"],
                    str((environment_dir / "rld.sqlite3").resolve()),
                )
                self.assertTrue((environment_dir / "rld.sqlite3").exists())

    def test_public_files_are_allowlisted(self):
        public_response = self.client.get("/", base_url=self.base_url)
        self.assertEqual(public_response.status_code, 200)
        public_response.close()
        for asset_path in (
            "/assets/icons/rld_logo.png",
            "/assets/icons/logout.png",
            "/assets/icons/clear.png",
            "/assets/icons/print.png",
            "/assets/icons/save.png",
            "/assets/icons/add.png",
            "/assets/icons/dark_mode.png",
            "/assets/icons/light_mode.png",
            "/assets/icons/images.png",
            "/assets/icons/search.png",
            "/assets/icons/sort.png",
            "/assets/icons/edit.png",
            "/assets/icons/delete.png",
            "/assets/icons/cancel.png",
            "/assets/icons/set_date.png",
            "/assets/images/header.png",
            "/assets/images/signature1.png",
            "/assets/images/signature2.png",
        ):
            with self.subTest(asset_path=asset_path):
                response = self.client.get(asset_path, base_url=self.base_url)
                self.assertEqual(response.status_code, 200)
                response.close()
        self.assertEqual(self.client.get("/app.py", base_url=self.base_url).status_code, 404)
        self.assertEqual(self.client.get("/data/units_price_list.json", base_url=self.base_url).status_code, 404)
        self.assertEqual(self.client.get("/instance/rld.sqlite3", base_url=self.base_url).status_code, 404)

    def test_login_errors_and_admin_password_change(self):
        missing_user = self.client.post(
            "/api/login", base_url=self.base_url, json={"username": "missing", "password": "anything"}
        )
        self.assertEqual(missing_user.status_code, 404)
        self.assertEqual(missing_user.json["code"], "username_not_found")

        wrong_password = self.client.post(
            "/api/login",
            base_url=self.base_url,
            json={"username": ADMIN_USERNAME, "password": "wrong"},
        )
        self.assertEqual(wrong_password.status_code, 401)
        self.assertEqual(wrong_password.json["code"], "wrong_password")

        self.assertEqual(self.login().status_code, 200)
        other_admin_session = self.app.test_client()
        self.assertEqual(self.login(other_admin_session).status_code, 200)
        changed = self.client.put(
            "/api/admin/password",
            base_url=self.base_url,
            json={"currentPassword": ADMIN_PASSWORD, "newPassword": "new-safe-password"},
        )
        self.assertEqual(changed.status_code, 200)
        self.assertEqual(self.client.get("/api/session", base_url=self.base_url).status_code, 401)
        self.assertEqual(
            other_admin_session.get("/api/session", base_url=self.base_url).status_code,
            401,
        )
        self.assertEqual(
            self.login(password="new-safe-password").status_code,
            200,
        )

    def test_admin_manages_accounts_and_regular_users_cannot(self):
        self.assertEqual(self.login().status_code, 200)
        created = self.client.post(
            "/api/users",
            base_url=self.base_url,
            json={"username": "mechanic", "password": "mechanic-password"},
        )
        self.assertEqual(created.status_code, 201)

        user_client = self.app.test_client()
        self.assertEqual(
            self.login(user_client, "mechanic", "mechanic-password").status_code,
            200,
        )
        self.assertEqual(user_client.get("/api/users", base_url=self.base_url).status_code, 403)
        reset = self.client.put(
            "/api/users/mechanic",
            base_url=self.base_url,
            json={"password": "reset-password"},
        )
        self.assertEqual(reset.status_code, 200)
        self.assertEqual(user_client.get("/api/session", base_url=self.base_url).status_code, 401)

    def test_authenticated_users_manage_shared_unit_price_list(self):
        self.assertEqual(
            self.client.get("/api/units", base_url=self.base_url).status_code,
            401,
        )
        self.assertEqual(self.login().status_code, 200)
        unit_data = {
            "maker": "Toyota",
            "unit": "Hilux 2.4 E",
            "transmission": "Automatic",
            "color": "Silver",
            "plateNo": "ABC-123",
            "yearModel": 2024,
            "fuelType": "Diesel",
            "sellingPrice": 1450000,
            "available": True,
        }
        created = self.client.post(
            "/api/units", base_url=self.base_url, json=unit_data
        )
        self.assertEqual(created.status_code, 201)
        saved = created.json["unit"]
        self.assertEqual(saved["maker"], "Toyota")
        self.assertEqual(saved["color"], "Silver")
        self.assertEqual(saved["sellingPrice"], 1450000)
        self.assertTrue(saved["available"])
        listed_units = self.client.get("/api/units", base_url=self.base_url).json["units"]
        self.assertIn(saved, listed_units)

        updated = self.client.put(
            f"/api/units/{saved['id']}",
            base_url=self.base_url,
            json={**unit_data, "unit": "Hilux 2.4 G", "available": False},
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json["unit"]["unit"], "Hilux 2.4 G")
        self.assertEqual(updated.json["unit"]["color"], "Silver")
        self.assertFalse(updated.json["unit"]["available"])
        invalid = self.client.put(
            f"/api/units/{saved['id']}",
            base_url=self.base_url,
            json={**unit_data, "available": "false"},
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(
            self.client.post(
                "/api/units",
                base_url=self.base_url,
                json={**unit_data, "maker": "Unknown"},
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.post(
                "/api/units",
                base_url=self.base_url,
                json={**unit_data, "yearModel": 2024.5},
            ).status_code,
            400,
        )

        unpriced = self.client.post(
            "/api/units",
            base_url=self.base_url,
            json={
                **unit_data,
                "maker": "Suzuki",
                "unit": "Unknown price",
                "color": "",
                "plateNo": "SUZ-001",
                "sellingPrice": None,
            },
        )
        self.assertEqual(unpriced.status_code, 201)
        self.assertIsNone(unpriced.json["unit"]["sellingPrice"])

        deleted = self.client.delete(
            f"/api/units/{saved['id']}", base_url=self.base_url
        )
        self.assertEqual(deleted.status_code, 200)
        remaining_ids = {
            unit["id"]
            for unit in self.client.get("/api/units", base_url=self.base_url).json["units"]
        }
        self.assertNotIn(saved["id"], remaining_ids)
        self.assertEqual(
            self.client.delete(
                f"/api/units/{saved['id']}", base_url=self.base_url
            ).status_code,
            404,
        )

    def test_bundled_price_list_includes_colors_and_expands_transmission_codes(self):
        self.assertEqual(self.login().status_code, 200)
        units = self.client.get("/api/units", base_url=self.base_url).json["units"]
        self.assertEqual(len(units), 29)
        city = next(unit for unit in units if unit["plateNo"] == "ABD 9854")
        self.assertEqual(city["unit"], "CITY")
        self.assertEqual(city["color"], "Modern Steel")
        self.assertEqual(city["transmission"], "Manual")
        dmax = next(unit for unit in units if unit["plateNo"] == "FTF 574")
        self.assertIsNone(dmax["sellingPrice"])

        restarted_app = create_app(
            database_path=Path(self.temporary_directory.name) / "test.sqlite3",
            secret_key="test-secret",
        )
        restarted_client = restarted_app.test_client()
        self.assertEqual(self.login(restarted_client).status_code, 200)
        self.assertEqual(
            len(restarted_client.get("/api/units", base_url=self.base_url).json["units"]),
            29,
        )

    def test_unit_images_are_private_capped_and_downloadable_in_a_zip(self):
        image_buffer = BytesIO()
        Image.new("RGB", (2, 2), "blue").save(image_buffer, format="PNG")
        image_bytes = image_buffer.getvalue()
        self.assertEqual(
            self.client.get(
                "/api/units/example/images", base_url=self.base_url
            ).status_code,
            401,
        )
        self.assertEqual(self.login().status_code, 200)
        unit = self.client.post(
            "/api/units",
            base_url=self.base_url,
            json={
                "maker": "Toyota",
                "unit": "Photo unit",
                "transmission": "Automatic",
                "plateNo": "PHOTO-1",
                "yearModel": 2024,
                "fuelType": "Diesel",
                "sellingPrice": None,
            },
        ).json["unit"]
        upload_response = self.client.post(
            f"/api/units/{unit['id']}/images",
            base_url=self.base_url,
            data={
                "images": [
                    (BytesIO(image_bytes), f"photo-{index}.png")
                    for index in range(20)
                ]
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(upload_response.status_code, 201)
        self.assertEqual(len(upload_response.json["images"]), 20)
        first_image = upload_response.json["images"][0]
        image_response = self.client.get(first_image["url"], base_url=self.base_url)
        self.assertEqual(image_response.data, image_bytes)
        image_response.close()
        too_many = self.client.post(
            f"/api/units/{unit['id']}/images",
            base_url=self.base_url,
            data={"images": (BytesIO(image_bytes), "extra.png")},
            content_type="multipart/form-data",
        )
        self.assertEqual(too_many.status_code, 400)
        invalid = self.client.post(
            f"/api/units/{unit['id']}/images",
            base_url=self.base_url,
            data={"images": (BytesIO(b"not an image"), "invalid.png")},
            content_type="multipart/form-data",
        )
        self.assertEqual(invalid.status_code, 400)
        second_image = upload_response.json["images"][1]
        download = self.client.get(
            f"/api/units/{unit['id']}/images-download"
            f"?imageIds={first_image['id']}&imageIds={second_image['id']}",
            base_url=self.base_url,
        )
        self.assertEqual(download.status_code, 200)
        with zipfile.ZipFile(BytesIO(download.data)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            self.assertEqual(
                [archive.read(name) for name in archive.namelist()],
                [image_bytes, image_bytes],
            )
        self.assertTrue(
            (Path(self.temporary_directory.name) / "unit-images").is_dir()
        )
        self.assertEqual(
            self.client.delete(
                f"/api/units/{unit['id']}", base_url=self.base_url
            ).status_code,
            200,
        )
        unit_image_directory = (
            Path(self.temporary_directory.name)
            / "unit-images"
            / sha256(unit["id"].encode("utf-8")).hexdigest()
        )
        self.assertFalse(unit_image_directory.exists())

    def test_existing_unit_database_migrates_color_and_unpriced_support(self):
        database_path = Path(self.temporary_directory.name) / "legacy.sqlite3"
        connection = sqlite3.connect(database_path)
        try:
            connection.execute(
                """
                CREATE TABLE units (
                    id TEXT PRIMARY KEY,
                    maker TEXT NOT NULL,
                    unit TEXT NOT NULL,
                    transmission TEXT NOT NULL,
                    plate_no TEXT NOT NULL,
                    year_model INTEGER NOT NULL,
                    fuel_type TEXT NOT NULL,
                    selling_price REAL NOT NULL CHECK (selling_price >= 0),
                    available INTEGER NOT NULL DEFAULT 1 CHECK (available IN (0, 1))
                )
                """
            )
            connection.execute(
                """
                INSERT INTO units
                    (id, maker, unit, transmission, plate_no, year_model,
                     fuel_type, selling_price, available)
                VALUES ('legacy-unit', 'Toyota', 'Legacy', 'Manual', 'LEGACY-1',
                    2010, 'Diesel', 100000, 1)
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated_app = create_app(database_path=database_path, secret_key="test-secret")
        migrated_client = migrated_app.test_client()
        self.assertEqual(self.login(migrated_client).status_code, 200)
        migrated_units = migrated_client.get(
            "/api/units", base_url=self.base_url
        ).json["units"]
        legacy = next(unit for unit in migrated_units if unit["id"] == "legacy-unit")
        self.assertEqual(legacy["color"], "")
        self.assertEqual(legacy["sellingPrice"], 100000)

    def test_estimate_sync_uses_latest_client_timestamp_and_tombstones(self):
        self.assertEqual(self.login().status_code, 200)
        estimate_id = "estimate-1"
        newest = {
            "id": estimate_id,
            "savedAt": "2026-01-01T00:00:00.000Z",
            "updatedAt": "2026-01-02T00:00:00.000Z",
            "data": {"clientName": "Latest version", "works": []},
            "deleted": False,
        }
        saved = self.client.put(
            f"/api/estimates/{estimate_id}", base_url=self.base_url, json=newest
        )
        self.assertEqual(saved.status_code, 200)
        self.assertTrue(saved.json["accepted"])

        stale = {**newest, "updatedAt": "2026-01-01T12:00:00.000Z", "data": {"clientName": "Old"}}
        stale_result = self.client.put(
            f"/api/estimates/{estimate_id}", base_url=self.base_url, json=stale
        )
        self.assertFalse(stale_result.json["accepted"])
        self.assertEqual(stale_result.json["estimate"]["data"]["clientName"], "Latest version")

        deleted = {**newest, "updatedAt": "2026-01-03T00:00:00.000Z", "deleted": True}
        tombstone = self.client.put(
            f"/api/estimates/{estimate_id}", base_url=self.base_url, json=deleted
        )
        self.assertTrue(tombstone.json["estimate"]["deleted"])
        history = self.client.get("/api/estimates", base_url=self.base_url)
        self.assertTrue(history.json["estimates"][0]["deleted"])

    def test_synced_estimate_is_saved_as_printable_excel(self):
        self.assertEqual(self.login().status_code, 200)
        response = self.client.put(
            "/api/estimates/excel-1",
            base_url=self.base_url,
            json={
                "id": "excel-1",
                "updatedAt": "2026-02-01T00:00:00.000Z",
                "data": {
                    "date": "2026-02-01",
                    "noteHiddenDamage": True,
                    "orgPrefix": "MGO Taal",
                    "makeModel": "Toyota Hilux",
                    "plateNo": "ABC-123",
                    "works": [
                        {"desc": "*parts", "labor": "", "parts": ""},
                        {"desc": "nut", "labor": 100, "parts": 0},
                        {"desc": "2. bolt", "labor": 200, "parts": 0},
                        {"desc": "*Tinsmith", "labor": "", "parts": ""},
                        {"desc": "side panel", "labor": 300, "parts": 0},
                        {"desc": "left panel", "labor": "", "parts": 0},
                    ],
                    "bundleLabor": True,
                    "bundleAllPrices": False,
                    "totLabor": 600,
                    "totParts": 0,
                    "downpayment": 100,
                    "totDue": 500,
                },
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json["accepted"])
        estimate_dir = (
            Path(self.temporary_directory.name)
            / "estimates"
            / sha256(b"excel-1").hexdigest()
        )
        workbook_path = estimate_dir / "MGO TAAL_TOYOTA HILUX_ABC-123.xlsx"
        self.assertTrue(workbook_path.exists())
        workbook = load_workbook(workbook_path)
        sheet = workbook["Job Estimate"]
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
        self.assertIn("1. NUT", values)
        self.assertEqual(sheet["A9"].value, "DATE: FEBRUARY 1, 2026")
        self.assertEqual(sheet["A23"].value, "NOTE: HIDDEN DAMAGE ARE NOT INCLUDED")
        self.assertTrue(sheet["A23"].font.bold)
        self.assertEqual(sheet["A23"].font.color.rgb, "00C00000")
        self.assertIn("2. BOLT", values)
        self.assertIn("1. SIDE PANEL", values)
        self.assertIn("*PARTS", values)
        self.assertIn("*TINSMITH", values)
        self.assertEqual(str(sheet.page_setup.paperSize), "14")
        self.assertEqual(sheet.column_dimensions["A"].width, 25)
        self.assertEqual(sheet.column_dimensions["B"].width, 25)
        self.assertEqual(sheet.column_dimensions["C"].width, 25)
        self.assertEqual(sheet.column_dimensions["D"].width, 25)
        self.assertEqual(sheet["A17"].alignment.horizontal, "left")
        self.assertEqual(sheet["C18"].alignment.horizontal, "right")
        self.assertEqual(sheet["A17"].border.left.style, "thin")
        self.assertEqual(sheet["D18"].border.right.style, "thin")
        self.assertIn("C21:C22", {str(merged) for merged in sheet.merged_cells.ranges})
        self.assertEqual(len(sheet._images), 3)
        self.assertEqual(round(sheet._images[1].anchor.ext.cy / 9525), 65)
        self.assertEqual(round(sheet._images[2].anchor.ext.cy / 9525), 65)
        self.assertEqual(
            sheet._images[1].anchor._from.colOff + sheet._images[1].anchor.ext.cx // 2,
            sheet._images[2].anchor._from.colOff + sheet._images[2].anchor.ext.cx // 2,
        )
        self.assertGreaterEqual(len(list(estimate_dir.glob("*.xlsx"))), 1)
        update = self.client.put(
            "/api/estimates/excel-1",
            base_url=self.base_url,
            json={
                "id": "excel-1",
                "updatedAt": "2026-02-02T00:00:00.000Z",
                "data": {
                    **response.json["estimate"]["data"],
                    "makeModel": "Toyota Prius",
                },
            },
        )
        self.assertEqual(update.status_code, 200)
        self.assertFalse(workbook_path.exists())
        self.assertTrue((estimate_dir / "MGO TAAL_TOYOTA PRIUS_ABC-123.xlsx").exists())
        self.assertEqual(
            self.client.get("/api/estimates/excel-1/pdf", base_url=self.base_url).status_code,
            404,
        )
        bundled = self.client.put(
            "/api/estimates/bundle-1",
            base_url=self.base_url,
            json={
                "id": "bundle-1",
                "updatedAt": "2026-02-03T00:00:00.000Z",
                "data": {
                    "works": [
                        {"desc": "*parts", "labor": "", "parts": ""},
                        {"desc": "bolt", "labor": "", "parts": 5000},
                        {"desc": "*mechanical", "labor": "", "parts": ""},
                        {"desc": "first repair", "labor": 10000, "parts": 1000},
                        {"desc": "second repair", "labor": 8000, "parts": 1000},
                    ],
                    "bundleLabor": True,
                    "bundleAllPrices": True,
                    "totLabor": 18000,
                    "totParts": 7000,
                },
            },
        )
        self.assertEqual(bundled.status_code, 200)
        bundled_dir = (
            Path(self.temporary_directory.name)
            / "estimates"
            / sha256(b"bundle-1").hexdigest()
        )
        bundled_workbook = load_workbook(bundled_dir / "PRIVATE_UNKNOWN VEHICLE_NO PLATE.xlsx")
        bundled_sheet = bundled_workbook["Job Estimate"]
        self.assertEqual(bundled_sheet["C17"].value, 25000)
        for row in (17, 18, 20, 21):
            if row != 17:
                self.assertIsNone(bundled_sheet[f"C{row}"].value)
            self.assertIsNone(bundled_sheet[f"D{row}"].value)
        self.assertIn("C17:D21", {str(merged) for merged in bundled_sheet.merged_cells.ranges})
        self.assertEqual(bundled_sheet["C17"].alignment.horizontal, "center")
        self.assertEqual(bundled_sheet["C22"].value, 25000)
        self.assertIn("C22:D22", {str(merged) for merged in bundled_sheet.merged_cells.ranges})


if __name__ == "__main__":
    unittest.main()
